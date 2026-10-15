"""
Project UNKNOWN (AETHER)

Document Loader v1.0

Reads supported research documents (PDF, DOCX, TXT),
extracts clean text with metadata, and stores structured JSON.
"""

from __future__ import annotations

import json
import time
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.rag.config import settings
from app.rag.exceptions import (
    DocumentExtractionError,
    DocumentNotFoundError,
    EmptyDocumentError,
    UnsupportedDocumentError,
)
from app.rag.logger import logger
from app.rag.utils import (
    clean_text,
    count_words,
    estimate_read_time,
    generate_hash,
    is_supported_document,
    list_supported_documents,
)


class DocumentLoader:
    """Enterprise-grade document ingestion engine."""

    # ----------------------------------------------------------
    # PDF Loader
    # ----------------------------------------------------------
    def _load_pdf(self, file_path: Path) -> list[dict]:
        reader = PdfReader(file_path)
        pages = []
        tess_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

        for page_number, page in enumerate(reader.pages, start=1):
            raw_text = page.extract_text() or ""
            text = clean_text(raw_text)

            # Fallback to OCR if page has no digital text layer (e.g. scanned papers / raster PDFs)
            if not text.strip() and hasattr(page, "images") and len(page.images) > 0:
                try:
                    import io
                    from PIL import Image
                    import pytesseract

                    if Path(tess_cmd).exists():
                        pytesseract.pytesseract.tesseract_cmd = tess_cmd

                    ocr_parts = []
                    for img_name, img_obj in page.images.items():
                        try:
                            pil_img = Image.open(io.BytesIO(img_obj.data))
                            # Skip tiny decorative icons/watermarks
                            if pil_img.width > 200 and pil_img.height > 200:
                                ocr_res = pytesseract.image_to_string(pil_img)
                                if ocr_res.strip():
                                    ocr_parts.append(ocr_res.strip())
                        except Exception as img_err:
                            logger.debug(f"Failed to OCR image {img_name}: {img_err}")

                    if ocr_parts:
                        text = clean_text("\n\n".join(ocr_parts))
                        logger.info(f"OCR extracted {len(text)} chars from page {page_number} of {file_path.name}")
                except Exception as ocr_err:
                    logger.warning(f"OCR fallback unavailable for page {page_number}: {ocr_err}")

            pages.append(
                {
                    "page_number": page_number,
                    "word_count": count_words(text),
                    "reading_time_minutes": estimate_read_time(text),
                    "page_hash": generate_hash(text),
                    "text": text,
                }
            )

        return pages

    # ----------------------------------------------------------
    # DOCX Loader
    # ----------------------------------------------------------
    def _load_docx(self, file_path: Path) -> list[dict]:
        try:
            document = DocxDocument(file_path)  # type: ignore
        except Exception as exc:
            raise DocumentExtractionError(f"Corrupted or invalid DOCX document: {exc}") from exc

        content_parts = []
        for paragraph in document.paragraphs:
            stripped = paragraph.text.strip()
            if stripped:
                content_parts.append(stripped)

        # Extract structured table content
        for table in document.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    content_parts.append(" | ".join(row_cells))

        text = clean_text("\n\n".join(content_parts))

        return [
            {
                "page_number": 1,
                "word_count": count_words(text),
                "reading_time_minutes": estimate_read_time(text),
                "page_hash": generate_hash(text),
                "text": text,
            }
        ]

    # ----------------------------------------------------------
    # TXT Loader
    # ----------------------------------------------------------
    def _load_txt(self, file_path: Path) -> list[dict]:
        raw_text = ""
        for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                raw_text = file_path.read_text(encoding=encoding)
                break
            except (UnicodeDecodeError, LookupError):
                continue

        if not raw_text:
            raw_text = file_path.read_text(encoding="utf-8", errors="replace")

        text = clean_text(raw_text)

        return [
            {
                "page_number": 1,
                "word_count": count_words(text),
                "reading_time_minutes": estimate_read_time(text),
                "page_hash": generate_hash(text),
                "text": text,
            }
        ]

    # ----------------------------------------------------------
    # Universal Loader
    # ----------------------------------------------------------
    def extract_document(self, file_path: Path) -> list[dict]:
        """
        Extract pages and metadata from any supported document.
        """

        if not file_path.exists():
            logger.error(f"Document not found: {file_path}")
            raise DocumentNotFoundError(file_path)

        if not is_supported_document(file_path):
            logger.error(f"Unsupported format: {file_path.suffix}")
            raise UnsupportedDocumentError(file_path.suffix)

        try:
            suffix = file_path.suffix.lower()

            if suffix == ".pdf":
                pages = self._load_pdf(file_path)

            elif suffix == ".docx":
                pages = self._load_docx(file_path)

            elif suffix == ".txt":
                pages = self._load_txt(file_path)

            else:
                raise UnsupportedDocumentError(suffix)

            full_text = " ".join(page["text"] for page in pages)

            if not full_text.strip():
                logger.error(f"Empty document: {file_path.name}")
                raise EmptyDocumentError(file_path.name)

            logger.info(f"Loaded {file_path.name} ({len(pages)} page(s))")

            return pages

        except EmptyDocumentError:
            raise

        except Exception as exc:  # noqa: BLE001
            logger.exception(f"Extraction failed: {file_path.name}")
            raise DocumentExtractionError(str(exc))

    # ----------------------------------------------------------
    # JSON Builder
    # ----------------------------------------------------------
    def create_document_json(self, file_path: Path) -> dict:
        pages = self.extract_document(file_path)

        full_text = " ".join(page["text"] for page in pages)

        return {
            "title": file_path.stem,
            "source_file": file_path.name,
            "document_type": file_path.suffix.lower(),
            "document_hash": generate_hash(full_text),
            "total_pages": len(pages),
            "total_words": count_words(full_text),
            "estimated_read_time": estimate_read_time(full_text),
            "pages": pages,
        }

    # ----------------------------------------------------------
    # Save JSON
    # ----------------------------------------------------------
    def save_document_json(self, document: dict) -> Path:
        output_path = settings.JSON_DIR / f"{document['title']}.json"

        with open(
            output_path,
            "w",
            encoding=settings.DEFAULT_ENCODING,
        ) as file:
            json.dump(document, file, indent=4, ensure_ascii=False)

        logger.info(f"Saved JSON → {output_path.name}")

        return output_path

    # ----------------------------------------------------------
    # Process Entire Dataset
    # ----------------------------------------------------------
    def process_all_documents(self) -> None:
        documents = list_supported_documents(settings.PAPERS_DIR)

        logger.info(f"Found {len(documents)} supported document(s).")
        print(f"Found {len(documents)} supported document(s).")

        start = time.perf_counter()

        for document_path in documents:
            logger.info(f"Processing {document_path.name}")

            document_json = self.create_document_json(document_path)
            self.save_document_json(document_json)

        elapsed = round(time.perf_counter() - start, 2)

        logger.info(f"Pipeline completed in {elapsed} seconds.")
        print(f"Pipeline completed in {elapsed} seconds.")


# ----------------------------------------------------------
# Entry Point
# ----------------------------------------------------------
if __name__ == "__main__":
    loader = DocumentLoader()
    loader.process_all_documents()
